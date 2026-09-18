"""Parse Wisdom Muyumbu students, staff, and photos.

Scoped to C:\\methode\\15 Wisdoms\\11.Wisdom Muyumbu only.
All student sheets are REB: Baby class, Middle Class, Top Class, P1–P6.
"""
from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl

ROOT = Path(r"C:\methode\15 Wisdoms\11.Wisdom Muyumbu")
WORKBOOK = ROOT / "students.xlsx"
STAFF_DOC = ROOT / "staffs.docx"
PHOTO_DIR = ROOT / "staffs photo"
OUTPUT = Path(r"C:\xampp7\htdocs\Xander-school\deploy\_wisdom_muyumbu_school_information.json")
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}

SHEET_TO_CLASS = {
    "BABY": "BABY CLASS",
    "BABY CLASS": "BABY CLASS",
    "MIDDLE": "MIDDLE CLASS",
    "MIDDLE CLASS": "MIDDLE CLASS",
    "TOP": "TOP CLASS",
    "TOP CLASS": "TOP CLASS",
    "P1": "P1",
    "P2": "P2",
    "P3": "P3",
    "P4": "P4",
    "P5": "P5",
    "P6": "P6",
}


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip(" \t:-.")


def split_name(value: str) -> tuple[str, str]:
    parts = clean(value).split(" ")
    if not parts or parts == [""]:
        return "", ""
    if len(parts) == 1:
        return parts[0], ""
    return parts[0], " ".join(parts[1:])


def clean_phone(value: object) -> str:
    text = clean(value).lstrip("'")
    digits = re.sub(r"[^0-9+]", "", text)
    return digits[:20]


def clean_email(value: object) -> str:
    text = clean(value).lower().replace(" ", "").replace("..", ".")
    if not text or text in {"-", "n/a", "na"}:
        return ""
    if "gmail" in text and "@" not in text:
        text = re.sub(r"gmail\.com$", "@gmail.com", text)
    if "@gmail" in text and not text.endswith("gmail.com"):
        text = re.sub(r"@gmail.*", "@gmail.com", text)
    return text[:100]


def normalize_position(value: str) -> str:
    compact = re.sub(r"[^A-Z]", "", clean(value).upper())
    if "HEAD" in compact:
        return "HEAD TEACHER"
    if "DOS" in compact:
        return "DOS"
    if "ACCOUNTANT" in compact:
        return "ACCOUNTANT"
    return "TEACHER"


def parse_students() -> list[dict[str, str]]:
    workbook = openpyxl.load_workbook(WORKBOOK, data_only=True, read_only=True)
    students: list[dict[str, str]] = []
    for sheet_name in workbook.sheetnames:
        class_label = SHEET_TO_CLASS.get(sheet_name.strip().upper())
        if class_label is None:
            continue
        ws = workbook[sheet_name]
        for row in ws.iter_rows(values_only=True):
            cells = [clean(cell) for cell in row]
            number = cells[0] if cells else ""
            if not number.isdigit():
                continue
            first = cells[1] if len(cells) > 1 else ""
            second = cells[2] if len(cells) > 2 else ""
            if first.upper() in {"FIRST NAME", "NAME", "NAMES"}:
                continue
            full_name = clean(f"{first} {second}")
            if not full_name:
                continue
            fname = first if first else split_name(full_name)[0]
            lname = second if second else split_name(full_name)[1]
            if not lname and " " in fname:
                fname, lname = split_name(fname)
            students.append(
                {
                    "class_label": class_label,
                    "sheet": sheet_name.strip(),
                    "full_name": full_name,
                    "fname": fname,
                    "lname": lname,
                }
            )
    workbook.close()
    return students


def cell_text(cell: ET.Element) -> str:
    return "".join((t.text or "") for t in cell.findall(".//w:t", NS)).strip()


def parse_staff() -> list[dict[str, str]]:
    staff: list[dict[str, str]] = []
    with zipfile.ZipFile(STAFF_DOC) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    for table in root.findall(".//w:tbl", NS):
        started = False
        for row in table.findall("w:tr", NS):
            cells = [clean(cell_text(cell)) for cell in row.findall("w:tc", NS)]
            joined = " ".join(cells).upper()
            if not started:
                if "NAME" in joined and "POSITION" in joined:
                    started = True
                continue
            number = re.sub(r"[^0-9]", "", cells[0] if cells else "")
            if not number:
                continue
            name = cells[1] if len(cells) > 1 else ""
            if not name:
                continue
            fname, lname = split_name(name)
            staff.append(
                {
                    "full_name": name,
                    "fname": fname,
                    "lname": lname,
                    "position": normalize_position(cells[2] if len(cells) > 2 else "TEACHER"),
                    "phone": clean_phone(cells[3] if len(cells) > 3 else ""),
                    "email": clean_email(cells[4] if len(cells) > 4 else ""),
                }
            )
    return staff


def parse_photos() -> list[dict[str, str]]:
    photos: list[dict[str, str]] = []
    if not PHOTO_DIR.is_dir():
        return photos
    for path in sorted(PHOTO_DIR.iterdir()):
        if path.suffix.lower() not in {".jpg", ".jpeg", ".png", ".webp"}:
            continue
        photos.append({"file": path.name, "name_hint": path.stem})
    return photos


def main() -> int:
    if not WORKBOOK.is_file():
        raise SystemExit(f"Missing {WORKBOOK}")
    if "MUYUMBU" not in str(ROOT).upper():
        raise SystemExit("Refusing to parse: folder is not Wisdom Muyumbu")
    students = parse_students()
    staff = parse_staff()
    photos = parse_photos()
    by_class: dict[str, int] = {}
    for student in students:
        by_class[student["class_label"]] = by_class.get(student["class_label"], 0) + 1
    head = next((row["full_name"] for row in staff if "HEAD" in row["position"].upper()), "")
    payload = {
        "school": {
            "name": "WISDOM SCHOOL MUYUMBU",
            "acronym": "WIS-MUY",
            "slogan": "",
            "academic_year": "2026-2027",
            "term": "Term I",
            "head_teacher": head,
        },
        "classes": students,
        "staff": staff,
        "photos": photos,
        "skipped": [],
    }
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(
        json.dumps(
            {
                "students": len(students),
                "staff": len(staff),
                "photos": len(photos),
                "by_class": by_class,
                "head_teacher": head,
            },
            indent=2,
        )
    )
    print("Staff:", len(staff))
    for row in staff:
        print(" -", row["full_name"], "|", row["position"], "|", row["phone"], "|", row["email"])
    print("Photos:")
    for photo in photos:
        print(" -", photo["file"])
    print("Wrote", OUTPUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
