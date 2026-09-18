"""Parse Wisdom Kiramuruzi students, staff, and captioned WhatsApp photos.

Scoped to C:\\methode\\15 Wisdoms\\4.KIRAMURUZI only.
Students are REB: Baby class, Middle Class, Top Class, P1–P6.
Photo filenames are WhatsApp timestamps; names are read from the caption on each picture.
"""
from __future__ import annotations

import json
import re
import shutil
from pathlib import Path

import openpyxl

ROOT = Path(r"C:\methode\15 Wisdoms\4.KIRAMURUZI")
OUTPUT = Path(r"C:\xampp7\htdocs\Xander-school\deploy\_wisdom_kiramuruzi_school_information.json")
PHOTO_DIR = Path(r"C:\xampp7\htdocs\Xander-school\deploy\_wisdom_kiramuruzi_staff_photos")

PHOTO_CAPTIONS = {
    "WhatsApp Image 2026-09-18 at 8.39.21 AM.jpeg": "Musonera Eric",
    "WhatsApp Image 2026-09-18 at 8.39.21 AM (1).jpeg": "Muteteri Regine",
    "WhatsApp Image 2026-09-18 at 8.39.22 AM.jpeg": "Emeritha",
    "WhatsApp Image 2026-09-18 at 8.39.22 AM (1).jpeg": "Phiona",
    "WhatsApp Image 2026-09-18 at 8.39.23 AM.jpeg": "Akimana Jean Baptiste",
    "WhatsApp Image 2026-09-18 at 8.39.23 AM (1).jpeg": "Himbaza Gad",
    "WhatsApp Image 2026-09-18 at 8.39.23 AM (2).jpeg": "Mutoni Naume",
    "WhatsApp Image 2026-09-18 at 8.39.24 AM.jpeg": "Habimana Emmanuel",
    "WhatsApp Image 2026-09-18 at 8.39.24 AM (1).jpeg": "Shafi Francisco",
    "WhatsApp Image 2026-09-18 at 8.39.24 AM (2).jpeg": "Ndikumana Moses",
    "WhatsApp Image 2026-09-18 at 8.39.25 AM.jpeg": "Niyomufasha Joseline",
    "WhatsApp Image 2026-09-18 at 8.39.25 AM (1).jpeg": "Manishimwe Elisa",
    "WhatsApp Image 2026-09-18 at 8.39.25 AM (2).jpeg": "Nshimiyimana Jean Bosco",
    "WhatsApp Image 2026-09-18 at 8.39.25 AM (3).jpeg": "Nyirahabineza Gaudance",
    "WhatsApp Image 2026-09-18 at 8.39.26 AM.jpeg": "Manirumva Aimable",
}

HEADER_TO_CLASS = {
    "BABY CLASS": "BABY CLASS",
    "MIDDLE CLASS": "MIDDLE CLASS",
    "TOP CLASS": "TOP CLASS",
    "P1": "P1",
    "P2": "P2",
    "P3": "P3",
    "P4": "P4",
    "P5": "P5",
    "P6": "P6",
    "PRIMARY SIX": "P6",
    "PRIMARY 6": "P6",
}


def find_workbook() -> Path:
    for path in ROOT.iterdir():
        if path.suffix.lower() == ".xlsx" and not path.name.startswith("~$"):
            return path
    raise SystemExit("Missing Kiramuruzi workbook")


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip(" \t:-.")


def split_name(value: str) -> tuple[str, str]:
    parts = clean(value).split(" ")
    if not parts or parts == [""]:
        return "", ""
    if len(parts) == 1:
        return parts[0], ""
    return parts[0], " ".join(parts[1:])


def compact(value: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", value.upper())


def name_tokens(value: str) -> set[str]:
    return {part for part in re.findall(r"[A-Z]+", value.upper()) if part not in {"TR", "H", "N", "M", "J"}}


def clean_phone(value: object) -> str:
    digits = re.sub(r"[^0-9]", "", clean(value).lstrip("'"))
    if len(digits) == 9 and digits.startswith("7"):
        digits = "0" + digits
    return digits[:20]


def clean_email(value: object) -> str:
    text = clean(value).lower().replace(" ", "").replace("..", ".")
    if not text or text in {"-", "n/a", "na"}:
        return ""
    if "gmail" in text and "@" not in text:
        text = re.sub(r"gmail\.com$", "@gmail.com", text)
    if "@gmail" in text and not text.endswith("gmail.com"):
        text = re.sub(r"@gmail.*", "@gmail.com", text)
    return text[:100]


def class_from_header(value: object) -> str:
    text = re.sub(r"[^A-Z0-9 ]", "", clean(value).upper())
    text = re.sub(r"\s+", " ", text).strip()
    if text in HEADER_TO_CLASS:
        return HEADER_TO_CLASS[text]
    compact_text = compact(text)
    for key, label in HEADER_TO_CLASS.items():
        if compact(key) == compact_text:
            return label
    return ""


def normalize_position(value: str) -> str:
    compact_text = re.sub(r"[^A-Z]", "", clean(value).upper())
    if "DEPUTY" in compact_text:
        return "DEPUTY HEAD TEACHER"
    if "HEAD" in compact_text:
        return "HEAD TEACHER"
    if "DOS" in compact_text:
        return "DOS"
    if "ACCOUNTANT" in compact_text or "BURSAR" in compact_text:
        return "ACCOUNTANT"
    if "COOK" in compact_text:
        return "COOK"
    if "CLEAN" in compact_text or "CLEARN" in compact_text:
        return "CLEANER"
    if "GUARD" in compact_text:
        return "GUARD"
    return "TEACHER"


def alias_key(value: str) -> str:
    key = compact(value)
    key = key.replace("EMNMANUEL", "EMMANUEL")
    key = key.replace("FRANSCISCO", "FRANCISCO")
    key = key.replace("AKIMABA", "AKIMANA")
    return key


def parse_students(workbook_path: Path) -> list[dict[str, str]]:
    workbook = openpyxl.load_workbook(workbook_path, data_only=True)
    ws = workbook["Students"]
    groups: list[tuple[int, str]] = []
    for col in range(1, (ws.max_column or 1) + 1):
        label = class_from_header(ws.cell(3, col).value)
        if label:
            groups.append((col, label))
    students: list[dict[str, str]] = []
    seen: set[tuple[str, str]] = set()
    for name_col, default_class in groups:
        current_class = default_class
        for row in range(4, (ws.max_row or 3) + 1):
            name = clean(ws.cell(row, name_col).value)
            switched = class_from_header(name)
            if switched:
                current_class = switched
                continue
            number = clean(ws.cell(row, name_col - 1).value)
            if not number.isdigit() or not name:
                continue
            key = (current_class, compact(name))
            if key in seen:
                continue
            seen.add(key)
            fname, lname = split_name(name)
            students.append(
                {
                    "class_label": current_class,
                    "sheet": "Students",
                    "full_name": name,
                    "fname": fname,
                    "lname": lname,
                    "sex": "U",
                }
            )
    workbook.close()
    return students


def parse_staff(workbook_path: Path) -> list[dict[str, str]]:
    workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    ws = workbook["Teachers"]
    staff: list[dict[str, str]] = []
    started = False
    for row in ws.iter_rows(values_only=True):
        cells = [clean(cell) for cell in (row or ())]
        joined = " ".join(cells).upper()
        if not started:
            if "NAMES" in joined and "POS" in joined:
                started = True
            continue
        if "SUPPORT" in joined:
            continue
        number = cells[0] if cells else ""
        name = cells[1] if len(cells) > 1 else ""
        if not number.isdigit() or not name:
            continue
        fname, lname = split_name(name)
        staff.append(
            {
                "full_name": name,
                "fname": fname,
                "lname": lname,
                "position": normalize_position(cells[2] if len(cells) > 2 else "TEACHER"),
                "phone": clean_phone(cells[3] if len(cells) > 3 else ""),
                "email": clean_email(cells[4] if len(cells) > 4 else ""),
                "photo": "",
            }
        )
    workbook.close()
    return staff


def best_staff(hint: str, staff: list[dict[str, str]], used: set[str]) -> dict[str, str] | None:
    want_key = alias_key(hint)
    want = name_tokens(hint)
    best = None
    best_score = 0.0
    for row in staff:
        if row["full_name"] in used:
            continue
        have_key = alias_key(row["full_name"])
        have = name_tokens(row["full_name"])
        score = 0.0
        if want_key and have_key and (want_key == have_key or want_key in have_key or have_key in want_key):
            score = 1.0
        else:
            overlap = len(want & have)
            if overlap:
                score = overlap / max(len(want), len(have), 1)
                if "GAD" in want and "GAD" in have:
                    score = max(score, 0.9)
                if "EMMANUEL" in alias_key(hint) and "EMMANUEL" in alias_key(row["full_name"]):
                    score = max(score, 0.85)
        if score > best_score:
            best_score = score
            best = row
    return best if best_score >= 0.28 else None


def extract_photos(staff: list[dict[str, str]]) -> list[dict[str, str]]:
    PHOTO_DIR.mkdir(parents=True, exist_ok=True)
    for old in PHOTO_DIR.iterdir():
        if old.is_file():
            old.unlink()
    photos: list[dict[str, str]] = []
    used: set[str] = set()
    for source_name, hint in PHOTO_CAPTIONS.items():
        src = ROOT / source_name
        if not src.is_file():
            continue
        person = best_staff(hint, staff, used)
        label = person["full_name"] if person else hint
        slug = re.sub(r"[^A-Za-z0-9]+", "_", label).strip("_")
        filename = f"{len(photos) + 1:02d}_{slug}.jpg"
        shutil.copyfile(src, PHOTO_DIR / filename)
        photos.append({"file": filename, "name_hint": hint})
        if person:
            person["photo"] = filename
            used.add(person["full_name"])
    return photos


def main() -> int:
    if "KIRAMURUZI" not in str(ROOT).upper():
        raise SystemExit("Refusing to parse: folder is not Wisdom Kiramuruzi")
    workbook_path = find_workbook()
    students = parse_students(workbook_path)
    staff = parse_staff(workbook_path)
    photos = extract_photos(staff)
    by_class: dict[str, int] = {}
    for student in students:
        by_class[student["class_label"]] = by_class.get(student["class_label"], 0) + 1
    head = next((row["full_name"] for row in staff if row["position"] == "HEAD TEACHER"), "")
    payload = {
        "school": {
            "name": "WISDOM SCHOOL KIRAMURUZI",
            "acronym": "WIS-KIR",
            "slogan": "",
            "academic_year": "2026-2027",
            "term": "Term I",
            "head_teacher": head,
        },
        "classes": students,
        "staff": staff,
        "photos": photos,
        "skipped": [],
    }
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(
        json.dumps(
            {
                "students": len(students),
                "staff": len(staff),
                "photos": len(photos),
                "by_class": by_class,
                "head_teacher": head,
            },
            indent=2,
        )
    )
    print("Staff:")
    for row in staff:
        print(" -", row["full_name"], "|", row["position"], "|", row["phone"], "|", row["email"], "|", row["photo"])
    print("Photos:")
    for photo in photos:
        print(" -", photo["file"], "=>", photo["name_hint"])
    unmatched = [row["full_name"] for row in staff if not row["photo"]]
    if unmatched:
        print("UNMATCHED STAFF", unmatched)
    print("Wrote", OUTPUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
