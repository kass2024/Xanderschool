"""Parse Wisdom Kabarore students, staff, and photos.

Scoped to C:\\methode\\15 Wisdoms\\8.WISDOM KABARORE only.
All student sheets are REB: Baby class, Middle Class, Top Class, P1–P6.
Staff come from staffs.txt plus named photo files.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(r"C:\methode\15 Wisdoms\8.WISDOM KABARORE")
STUDENTS_XLSX = ROOT / "Students List.xlsx"
STAFF_TXT = ROOT / "staffs.txt"
PHOTO_DIR = ROOT / "Photos"
OUTPUT = Path(r"C:\xampp7\htdocs\Xander-school\deploy\_wisdom_kabarore_school_information.json")

SHEET_TO_CLASS = {
    "BABY": "BABY CLASS",
    "BABY CLASS": "BABY CLASS",
    "MIDDLE": "MIDDLE CLASS",
    "MIDDLE CLASS": "MIDDLE CLASS",
    "TOP": "TOP CLASS",
    "TOP CLASS": "TOP CLASS",
    "PRIMARY ONE": "P1",
    "PRIMARY 1": "P1",
    "P1": "P1",
    "PRIMARY TWO": "P2",
    "PRIMARY 2": "P2",
    "P2": "P2",
    "PRIMARY THREE": "P3",
    "PRIMARY 3": "P3",
    "P3": "P3",
    "PRIMARY FOUR": "P4",
    "PRIMARY 4": "P4",
    "P4": "P4",
    "PRIMARY FIVE": "P5",
    "PRIMARY 5": "P5",
    "P5": "P5",
    "PRIMARY SIX": "P6",
    "PRIMARY 6": "P6",
    "P6": "P6",
}

HEADER_WORDS = {
    "NAMES",
    "NAME",
    "FIRST NAME",
    "SECOND NAME",
    "NR",
    "N/B",
    "AMOUNT DUE",
    "AM0UNT DUE",
    "SCHOOL FEES",
    "BREAK F",
    "FOOD",
    "REPORT CARD",
    "AMOUNT PAID",
    "BALANCE",
    "REMEDIAL",
}

SKIP_NAME = re.compile(
    r"prepared|head teacher|innocent benurwanda|baby class|top class|middle class|primary",
    re.I,
)


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip(" \t:-.")


def split_name(value: str) -> tuple[str, str]:
    parts = clean(value).split(" ")
    if not parts or parts == [""]:
        return "", ""
    if len(parts) == 1:
        return parts[0], ""
    return parts[0], " ".join(parts[1:])


def clean_phone(value: object) -> str:
    text = clean(value).lstrip("'")
    digits = re.sub(r"[^0-9+]", "", text)
    return digits[:20]


def clean_email(value: object) -> str:
    text = clean(value).lower().replace(" ", "").replace("..", ".")
    if not text or text in {"-", "n/a", "na"}:
        return ""
    text = re.sub(r"@(\d+)gmail\.com$", r"\1@gmail.com", text)
    if text.endswith("@ail.com"):
        text = text.replace("@ail.com", "@gmail.com")
    if "gmail" in text and "@" not in text:
        text = re.sub(r"gmail\.com$", "@gmail.com", text)
    if "@gmail" in text and not text.endswith("gmail.com"):
        text = re.sub(r"@gmail.*", "@gmail.com", text)
    return text[:100]


def normalize_position(value: str) -> str:
    compact = re.sub(r"[^A-Z]", "", clean(value).upper())
    if "HEAD" in compact:
        return "HEAD TEACHER"
    if "DOS" in compact or "DIRECTOROFSTUDIES" in compact:
        return "DOS"
    if "ACCOUNTANT" in compact or "BURSAR" in compact:
        return "ACCOUNTANT"
    return "TEACHER"


def photo_hint(stem: str) -> str:
    text = stem.split(",")[0]
    text = re.sub(r"^TR\s*", "", text, flags=re.I)
    text = re.split(r"IMG[_-]?\d", text, maxsplit=1, flags=re.I)[0]
    text = re.sub(r"[^A-Za-z ]", " ", text)
    return clean(text)


def is_int(value: object) -> bool:
    text = clean(value)
    return text.isdigit() and 1 <= int(text) <= 99


def is_name_text(value: object) -> bool:
    text = clean(value)
    if not text:
        return False
    upper = text.upper()
    if upper in HEADER_WORDS or upper.replace(".", "") in HEADER_WORDS:
        return False
    if SKIP_NAME.search(text):
        return False
    if not re.search(r"[A-Za-z]", text):
        return False
    if re.fullmatch(r"[\d.,+\-]+", text):
        return False
    if text.upper() in {"BK", "S", "+", "H"}:
        return False
    return True


def join_name_parts(parts: list[str]) -> str:
    tokens: list[str] = []
    for part in parts:
        for token in clean(part).split(" "):
            if not token:
                continue
            if tokens and tokens[-1].upper() == token.upper():
                continue
            tokens.append(token)
    return " ".join(tokens)


def parse_students() -> list[dict[str, str]]:
    workbook = openpyxl.load_workbook(STUDENTS_XLSX, data_only=True, read_only=True)
    students: list[dict[str, str]] = []
    seen: set[tuple[str, str]] = set()
    for sheet_name in workbook.sheetnames:
        class_label = SHEET_TO_CLASS.get(sheet_name.strip().upper())
        if class_label is None:
            continue
        ws = workbook[sheet_name]
        for row in ws.iter_rows(values_only=True):
            cells = list(row)
            number_idx = next((i for i, cell in enumerate(cells) if is_int(cell)), None)
            if number_idx is None:
                continue
            name_cells = []
            for cell in cells[number_idx + 1 :]:
                if is_name_text(cell):
                    name_cells.append(clean(cell))
                elif name_cells:
                    break
            name = join_name_parts(name_cells[:2] if len(name_cells) >= 2 else name_cells)
            if not name:
                continue
            key = (class_label, re.sub(r"[^A-Z0-9]", "", name.upper()))
            if key in seen:
                continue
            seen.add(key)
            fname, lname = split_name(name)
            students.append(
                {
                    "class_label": class_label,
                    "sheet": sheet_name.strip(),
                    "full_name": name,
                    "fname": fname,
                    "lname": lname,
                    "sex": "U",
                }
            )
    workbook.close()
    return students


def parse_staff_txt() -> dict[str, dict[str, str]]:
    text = STAFF_TXT.read_text(encoding="utf-8", errors="replace")
    blocks = re.split(r"\n\s*\n", text.split("STAFF.", 1)[-1])
    by_key: dict[str, dict[str, str]] = {}
    for block in blocks:
        lines = [clean(line) for line in block.splitlines() if clean(line)]
        if not lines:
            continue
        first = re.sub(r"^\d+\s*:\s*", "", lines[0])
        name_part, _, rest = first.partition(",")
        name = clean(name_part)
        position = normalize_position(rest or " ".join(lines[1:3]))
        phone = ""
        email = ""
        joined = " ".join(lines)
        phone_m = re.search(r"(0\d{8,12})", joined.replace(" ", ""))
        if phone_m:
            phone = phone_m.group(1)
        email_m = re.search(r"([A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,})", joined)
        if email_m:
            email = clean_email(email_m.group(1))
        if not name or len(name.split()) < 2:
            # DoS block starts with position on previous numbered item
            if "yambabariye" in joined.lower() or "aroan" in joined.lower():
                name = "Aroan Yambabariye"
                position = "DOS"
            else:
                continue
        fname, lname = split_name(name)
        key = re.sub(r"[^A-Z0-9]", "", name.upper())
        by_key[key] = {
            "full_name": name,
            "fname": fname,
            "lname": lname,
            "position": position,
            "phone": clean_phone(phone),
            "email": email,
        }
    return by_key


def parse_photos() -> tuple[list[dict[str, str]], dict[str, dict[str, str]]]:
    photos: list[dict[str, str]] = []
    from_photos: dict[str, dict[str, str]] = {}
    if not PHOTO_DIR.is_dir():
        return photos, from_photos
    for path in sorted(PHOTO_DIR.iterdir()):
        if path.suffix.lower() not in {".jpg", ".jpeg", ".png", ".webp"}:
            continue
        hint = photo_hint(path.stem)
        photos.append({"file": path.name, "name_hint": hint})
        if not hint or hint.upper().startswith("WHATSAPP"):
            continue
        position = "TEACHER"
        if "," in path.stem:
            position = normalize_position(path.stem.split(",", 1)[1])
        fname, lname = split_name(hint)
        key = re.sub(r"[^A-Z0-9]", "", hint.upper())
        from_photos[key] = {
            "full_name": hint,
            "fname": fname,
            "lname": lname,
            "position": position,
            "phone": "",
            "email": "",
        }
    return photos, from_photos


def merge_staff(from_txt: dict[str, dict[str, str]], from_photos: dict[str, dict[str, str]]) -> list[dict[str, str]]:
    merged: dict[str, dict[str, str]] = {}

    def name_tokens(value: str) -> set[str]:
        return set(re.findall(r"[A-Z]+", value.upper()))

    def find_key(name: str) -> str | None:
        want = name_tokens(name)
        compact = re.sub(r"[^A-Z0-9]", "", name.upper())
        if compact in merged:
            return compact
        for key, row in merged.items():
            have = name_tokens(row["full_name"])
            if want and have and (want <= have or have <= want or len(want & have) >= 2):
                return key
        return None

    for source in (from_photos, from_txt):
        for row in source.values():
            key = find_key(row["full_name"]) or re.sub(r"[^A-Z0-9]", "", row["full_name"].upper())
            if key not in merged:
                merged[key] = dict(row)
                continue
            existing = merged[key]
            if row["position"] in {"HEAD TEACHER", "DOS", "ACCOUNTANT"}:
                existing["position"] = row["position"]
            if row.get("phone"):
                existing["phone"] = row["phone"]
            if row.get("email"):
                existing["email"] = row["email"]
            # Prefer staffs.txt name order when it has a phone/email
            if row.get("phone") or row.get("email"):
                existing["full_name"] = row["full_name"]
                existing["fname"] = row["fname"]
                existing["lname"] = row["lname"]

    rank = {"HEAD TEACHER": 0, "DOS": 1, "ACCOUNTANT": 2, "TEACHER": 3}
    staff = list(merged.values())
    staff.sort(key=lambda row: (rank.get(row["position"], 9), row["full_name"].upper()))
    return staff


def main() -> int:
    if not STUDENTS_XLSX.is_file() or not STAFF_TXT.is_file():
        raise SystemExit("Missing Kabarore student workbook or staffs.txt")
    if "KABARORE" not in str(ROOT).upper():
        raise SystemExit("Refusing to parse: folder is not Wisdom Kabarore")
    students = parse_students()
    photos, from_photos = parse_photos()
    staff = merge_staff(parse_staff_txt(), from_photos)
    by_class: dict[str, int] = {}
    for student in students:
        by_class[student["class_label"]] = by_class.get(student["class_label"], 0) + 1
    head = next((row["full_name"] for row in staff if row["position"] == "HEAD TEACHER"), "")
    payload = {
        "school": {
            "name": "WISDOM SCHOOL KABARORE",
            "acronym": "WIS-KAB",
            "slogan": "",
            "academic_year": "2026-2027",
            "term": "Term I",
            "head_teacher": head,
        },
        "classes": students,
        "staff": staff,
        "photos": photos,
        "skipped": [],
    }
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(
        json.dumps(
            {
                "students": len(students),
                "staff": len(staff),
                "photos": len(photos),
                "by_class": by_class,
                "head_teacher": head,
            },
            indent=2,
        )
    )
    print("Staff:")
    for row in staff:
        print(" -", row["full_name"], "|", row["position"], "|", row["phone"], "|", row["email"])
    print("Photos:")
    for photo in photos:
        print(" -", photo["file"], "=>", photo["name_hint"])
    print("Wrote", OUTPUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
