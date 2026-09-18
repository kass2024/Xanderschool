"""Parse Wisdom School Susa student Excel + staff Word + photos.

Scoped to C:\\methode\\15 Wisdoms\\6.WISDOM SCHOOL SUSA only.
The P3 sheet is labeled WISDOM SCHOOL NYABIHU and is recorded as skipped
so it is never mixed into Susa.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(r"C:\methode\15 Wisdoms\6.WISDOM SCHOOL SUSA")
STUDENTS = ROOT / "STUDENTS SUSA.xlsx"
STAFF_DOC = ROOT / "STAFFS.docx"
PHOTO_DIR = ROOT / "staff pictures"
OUTPUT = Path(r"C:\xampp7\htdocs\Xander-school\deploy\_wisdom_susa_school_information.json")

SHEET_TO_CLASS = {
    "BABY": "BABY CLASS",
    "MIDDLE": "MIDDLE CLASS",
    "TOP": "TOP CLASS",
    "P1": "P1",
    "P2": "P2",
    "P3": "P3",
    "P4": "P4",
    "P5": "P5",
    "P6": "P6",
}


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def split_name(value: str) -> tuple[str, str]:
    parts = clean(value).split(" ")
    if not parts or parts == [""]:
        return "", ""
    if len(parts) == 1:
        return parts[0], ""
    return parts[0], " ".join(parts[1:])


def compact(value: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", value.upper())


def parse_students() -> tuple[list[dict[str, str]], list[dict]]:
    workbook = openpyxl.load_workbook(STUDENTS, data_only=True, read_only=True)
    students: list[dict[str, str]] = []
    skipped: list[dict] = []
    for sheet_name in workbook.sheetnames:
        class_label = SHEET_TO_CLASS.get(sheet_name.upper(), sheet_name.upper())
        ws = workbook[sheet_name]
        rows = [list(row) for row in ws.iter_rows(values_only=True)]
        header_blob = " ".join(clean(cell) for row in rows[:8] for cell in row[:4])
        foreign = "NYABIHU" in header_blob.upper() and "SUSA" not in header_blob.upper()
        started = False
        sheet_students: list[dict[str, str]] = []
        for row in rows:
            cells = [clean(cell) for cell in row]
            joined = " ".join(cells).upper()
            if not started:
                if "STUDENT NAME" in joined:
                    started = True
                continue
            number = cells[0] if cells else ""
            name = cells[1] if len(cells) > 1 else ""
            if not number.isdigit():
                if sheet_students:
                    break
                continue
            if not name:
                continue
            fname, lname = split_name(name)
            sheet_students.append(
                {
                    "class_label": class_label,
                    "sheet": sheet_name,
                    "full_name": name,
                    "fname": fname,
                    "lname": lname,
                }
            )
        if foreign:
            skipped.append(
                {
                    "sheet": sheet_name,
                    "class_label": class_label,
                    "reason": "Sheet header is WISDOM SCHOOL NYABIHU — not imported into Susa",
                    "header": header_blob,
                    "students": sheet_students,
                }
            )
        else:
            students.extend(sheet_students)
    workbook.close()
    return students, skipped


def parse_staff_doc() -> tuple[dict[str, str], list[dict[str, str]]]:
    import zipfile

    with zipfile.ZipFile(STAFF_DOC) as zipped:
        xml = zipped.read("word/document.xml").decode("utf-8", "replace")
    text = re.sub(r"</w:p>", "\n", xml)
    text = re.sub(r"<[^>]+>", "", text)
    text = text.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
    lines = [clean(line) for line in text.splitlines() if clean(line)]

    school: dict[str, str] = {
        "name": "WISDOM SCHOOL SUSA",
        "acronym": "WIS-SUS",
        "slogan": "",
        "academic_year": "2026-2027",
        "term": "Term I",
        "head_teacher": "",
        "city": "Susa",
        "country": "Rwanda",
    }
    for line in lines:
        upper = line.upper()
        if "FEARING GOD" in upper:
            school["slogan"] = line
        if "HEAD TEACHER:" in upper:
            school["head_teacher"] = clean(line.split(":", 1)[-1])

    staff: list[dict[str, str]] = []
    i = 0
    while i < len(lines):
        if re.fullmatch(r"\d+", lines[i]) and i + 1 < len(lines):
            full_name = lines[i + 1]
            if compact(full_name) in {"NAMES", "POSITION", "PHONENUMBER", "EMAILADDRESSES"}:
                i += 1
                continue
            position = lines[i + 2] if i + 2 < len(lines) else ""
            phone = lines[i + 3] if i + 3 < len(lines) else ""
            email = lines[i + 4] if i + 4 < len(lines) else ""
            if not re.search(r"\d", phone):
                phone = ""
            if "@" not in email:
                email = ""
            fname, lname = split_name(full_name)
            staff.append(
                {
                    "full_name": full_name,
                    "fname": fname,
                    "lname": lname,
                    "position": position,
                    "phone": re.sub(r"[^\d+]", "", phone),
                    "email": email.lower(),
                }
            )
            i += 5
            continue
        i += 1
    return school, staff


def parse_photos() -> list[dict[str, str]]:
    photos = []
    if not PHOTO_DIR.is_dir():
        return photos
    for path in sorted(PHOTO_DIR.iterdir()):
        if path.suffix.lower() not in {".jpg", ".jpeg", ".png", ".webp"}:
            continue
        photos.append(
            {
                "file": path.name,
                "path": str(path),
                "name_hint": path.stem,
            }
        )
    return photos


def main() -> int:
    students, skipped = parse_students()
    school, staff = parse_staff_doc()
    photos = parse_photos()
    payload = {
        "school": school,
        "staff": staff,
        "photos": photos,
        "classes": students,
        "skipped": skipped,
        "counts": {
            "students": len(students),
            "staff": len(staff),
            "photos": len(photos),
            "skipped_students": sum(len(item.get("students") or []) for item in skipped),
            "by_class": {},
        },
    }
    by_class: dict[str, int] = {}
    for row in students:
        label = row["class_label"]
        by_class[label] = by_class.get(label, 0) + 1
    payload["counts"]["by_class"] = by_class
    OUTPUT.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(payload["counts"], indent=2))
    print("Skipped:", json.dumps(skipped, indent=2, ensure_ascii=False)[:800])
    print("Staff:", len(staff))
    for person in staff:
        print(" -", person["full_name"], "|", person["position"], "|", person["phone"])
    print("Wrote", OUTPUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
