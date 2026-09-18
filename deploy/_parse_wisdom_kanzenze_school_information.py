"""Parse Wisdom Kanzenze students, staff, and Excel-embedded photos.

Scoped to C:\\methode\\15 Wisdoms\\9.Wisdom KANZENZE only.
All student sheets are REB: Baby class, Middle Class, Top Class, P1–P6.
Staff photos are taken from the STAFF sheet PICTURE column.
"""
from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl

ROOT = Path(r"C:\methode\15 Wisdoms\9.Wisdom KANZENZE")
OUTPUT = Path(r"C:\xampp7\htdocs\Xander-school\deploy\_wisdom_kanzenze_school_information.json")
PHOTO_DIR = Path(r"C:\xampp7\htdocs\Xander-school\deploy\_wisdom_kanzenze_staff_photos")
NS = {
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
}

SHEET_TO_CLASS = {
    "N1": "BABY CLASS",
    "BABY": "BABY CLASS",
    "BABY CLASS": "BABY CLASS",
    "N2": "MIDDLE CLASS",
    "MIDDLE": "MIDDLE CLASS",
    "MIDDLE CLASS": "MIDDLE CLASS",
    "N3": "TOP CLASS",
    "TOP": "TOP CLASS",
    "TOP CLASS": "TOP CLASS",
    "P1": "P1",
    "P2": "P2",
    "P3": "P3",
    "P4": "P4",
    "P5": "P5",
    "P6": "P6",
}


def find_workbook() -> Path:
    for path in ROOT.iterdir():
        if path.suffix.lower() == ".xlsx" and not path.name.startswith("~$"):
            return path
    raise SystemExit("Missing Kanzenze workbook")


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip(" \t:-.")


def split_name(value: str) -> tuple[str, str]:
    parts = clean(value).split(" ")
    if not parts or parts == [""]:
        return "", ""
    if len(parts) == 1:
        return parts[0], ""
    return parts[0], " ".join(parts[1:])


def clean_phone(value: object) -> str:
    digits = re.sub(r"[^0-9]", "", clean(value).lstrip("'"))
    if not digits:
        return ""
    if len(digits) == 9 and digits.startswith("7"):
        digits = "0" + digits
    return digits[:20]


def clean_email(value: object) -> str:
    text = clean(value).lower().replace(" ", "").replace("..", ".")
    if not text or text in {"-", "n/a", "na"}:
        return ""
    if "gmail" in text and "@" not in text:
        text = re.sub(r"gmail\.com$", "@gmail.com", text)
    if "@gmail" in text and not text.endswith("gmail.com"):
        text = re.sub(r"@gmail.*", "@gmail.com", text)
    return text[:100]


def clean_sex(value: object) -> str:
    text = clean(value).upper()
    if text in {"M", "MALE", "BOY"}:
        return "M"
    if text in {"F", "FEMALE", "GIRL"}:
        return "F"
    return "U"


def normalize_position(value: str) -> str:
    compact = re.sub(r"[^A-Z]", "", clean(value).upper())
    if "HEAD" in compact:
        return "HEAD TEACHER"
    if "DOS" in compact:
        return "DOS"
    if "ACCOUNTANT" in compact or "BURSAR" in compact:
        return "ACCOUNTANT"
    return "TEACHER"


def skip_name(value: str) -> bool:
    upper = value.upper()
    if not value or not re.search(r"[A-Za-z]", value):
        return True
    if upper in {"NO", "N0", "NAMES", "NAME", "SEX", "SES"}:
        return True
    if "WISDOM SCHOOL" in upper or "FEARING GOD" in upper or "PRIMARY" in upper:
        return True
    return False


def parse_students(workbook_path: Path) -> list[dict[str, str]]:
    workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    students: list[dict[str, str]] = []
    seen: set[tuple[str, str]] = set()
    for sheet_name in workbook.sheetnames:
        class_label = SHEET_TO_CLASS.get(sheet_name.strip().upper())
        if class_label is None:
            continue
        ws = workbook[sheet_name]
        for row in ws.iter_rows(values_only=True):
            cells = [clean(cell) for cell in (row or ())]
            number = cells[0] if cells else ""
            name = cells[1] if len(cells) > 1 else ""
            sex = clean_sex(cells[2] if len(cells) > 2 else "")
            if not number.isdigit() or skip_name(name):
                continue
            key = (class_label, re.sub(r"[^A-Z0-9]", "", name.upper()))
            if key in seen:
                continue
            seen.add(key)
            fname, lname = split_name(name)
            students.append(
                {
                    "class_label": class_label,
                    "sheet": sheet_name.strip(),
                    "full_name": name,
                    "fname": fname,
                    "lname": lname,
                    "sex": sex,
                }
            )
    workbook.close()
    return students


def parse_staff(workbook_path: Path) -> list[dict[str, str]]:
    workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    if "STAFF" not in workbook.sheetnames:
        workbook.close()
        raise SystemExit("Missing STAFF sheet")
    ws = workbook["STAFF"]
    staff: list[dict[str, str]] = []
    started = False
    for row in ws.iter_rows(values_only=True):
        cells = [clean(cell) for cell in (row or ())]
        joined = " ".join(cells).upper()
        if not started:
            if "NAMES" in joined and "POSITION" in joined:
                started = True
            continue
        number = cells[0] if cells else ""
        name = cells[1] if len(cells) > 1 else ""
        if not number.isdigit() or skip_name(name):
            continue
        fname, lname = split_name(name)
        staff.append(
            {
                "full_name": name,
                "fname": fname,
                "lname": lname,
                "position": normalize_position(cells[2] if len(cells) > 2 else "TEACHER"),
                "phone": clean_phone(cells[3] if len(cells) > 3 else ""),
                "email": clean_email(cells[4] if len(cells) > 4 else ""),
                "photo": "",
            }
        )
    workbook.close()
    return staff


def extract_staff_photos(workbook_path: Path, staff: list[dict[str, str]]) -> list[dict[str, str]]:
    PHOTO_DIR.mkdir(parents=True, exist_ok=True)
    for old in PHOTO_DIR.iterdir():
        if old.is_file():
            old.unlink()
    with zipfile.ZipFile(workbook_path) as zf:
        drawing = zf.read("xl/drawings/drawing1.xml")
        rels = zf.read("xl/drawings/_rels/drawing1.xml.rels")
        rel_root = ET.fromstring(rels)
        rid_to_media = {}
        for rel in rel_root:
            rid = rel.attrib.get("Id", "")
            target = rel.attrib.get("Target", "")
            if rid and target:
                rid_to_media[rid] = "xl/drawings/" + target
                rid_to_media[rid] = str(Path("xl/drawings") / target).replace("\\", "/")
        pics: list[tuple[int, str]] = []
        draw_root = ET.fromstring(drawing)
        for pic in draw_root.findall(".//xdr:pic", NS):
            blip = pic.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip")
            off = pic.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}off")
            if blip is None:
                continue
            embed = blip.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed", "")
            y = int(off.attrib.get("y", "0")) if off is not None else 0
            media = rid_to_media.get(embed, "")
            if media:
                pics.append((y, media))
        pics.sort(key=lambda item: item[0])
        photos: list[dict[str, str]] = []
        for index, (_y, media) in enumerate(pics):
            src = media
            if src not in zf.namelist() and src.startswith("xl/drawings/"):
                src = str(Path("xl") / Path(src).name).replace("\\", "/")
            # drawing rel target is ../media/imageN.jpeg
            candidate = Path("xl/media") / Path(media).name
            if str(candidate).replace("\\", "/") in zf.namelist():
                src = str(candidate).replace("\\", "/")
            person = staff[index] if index < len(staff) else None
            hint = person["full_name"] if person else f"staff {index + 1}"
            slug = re.sub(r"[^A-Za-z0-9]+", "_", hint).strip("_")
            filename = f"{index + 1:02d}_{slug}.jpg"
            dest = PHOTO_DIR / filename
            dest.write_bytes(zf.read(src))
            photos.append({"file": filename, "name_hint": hint})
            if person:
                person["photo"] = filename
        return photos


def main() -> int:
    if "KANZENZE" not in str(ROOT).upper():
        raise SystemExit("Refusing to parse: folder is not Wisdom Kanzenze")
    workbook_path = find_workbook()
    students = parse_students(workbook_path)
    staff = parse_staff(workbook_path)
    photos = extract_staff_photos(workbook_path, staff)
    by_class: dict[str, int] = {}
    for student in students:
        by_class[student["class_label"]] = by_class.get(student["class_label"], 0) + 1
    head = next((row["full_name"] for row in staff if row["position"] == "HEAD TEACHER"), "")
    payload = {
        "school": {
            "name": "WISDOM SCHOOL KANZENZE",
            "acronym": "WIS-KAN",
            "slogan": "FEARING GOD IS KNOWLEDGE",
            "academic_year": "2026-2027",
            "term": "Term I",
            "head_teacher": head,
        },
        "classes": students,
        "staff": staff,
        "photos": photos,
        "skipped": [],
    }
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(
        json.dumps(
            {
                "students": len(students),
                "staff": len(staff),
                "photos": len(photos),
                "by_class": by_class,
                "head_teacher": head,
            },
            indent=2,
        )
    )
    print("Staff:")
    for row in staff:
        print(" -", row["full_name"], "|", row["position"], "|", row["phone"], "|", row["email"], "|", row["photo"])
    print("Photos:")
    for photo in photos:
        print(" -", photo["file"], "=>", photo["name_hint"])
    print("Wrote", OUTPUT)
    print("Photos dir", PHOTO_DIR)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
