"""Parse Wisdom Burera students, staff, and photos.

Scoped to C:\\methode\\15 Wisdoms\\3.WISDOM SCHOOL BURERA only.
All student sheets are REB: Baby class, Middle Class, Top Class, P1–P6.
"""
from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl

ROOT = Path(r"C:\methode\15 Wisdoms\3.WISDOM SCHOOL BURERA")
WORKBOOK = ROOT / "LEARNERS OF WUSDOM SCHOOL BURERA WHO WILL PARTICIPATE IN THE PEN PAL PROGRAM.xlsx"
STAFF_DOC = ROOT / "STAFF BURERA.docx"
PHOTO_DIR = ROOT / "STAFF PHOTO"
OUTPUT = Path(r"C:\xampp7\htdocs\Xander-school\deploy\_wisdom_burera_school_information.json")
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}

SHEET_TO_CLASS = {
    "BABY CLASS": "BABY CLASS",
    "BABY": "BABY CLASS",
    "MIDDEL CLASS": "MIDDLE CLASS",
    "MIDDLE CLASS": "MIDDLE CLASS",
    "MIDDLE": "MIDDLE CLASS",
    "TOP CLASS": "TOP CLASS",
    "TOP": "TOP CLASS",
    "P1": "P1",
    "P2": "P2",
    "P3": "P3",
    "P4": "P4",
    "P5": "P5",
    "P6": "P6",
}


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip(" \t:-")


def split_name(value: str) -> tuple[str, str]:
    parts = clean(value).split(" ")
    if not parts or parts == [""]:
        return "", ""
    if len(parts) == 1:
        return parts[0], ""
    return parts[0], " ".join(parts[1:])


def clean_phone(value: object) -> str:
    text = clean(value).lstrip("'")
    digits = re.sub(r"[^0-9+]", "", text)
    return digits[:20]


def clean_email(value: object) -> str:
    text = clean(value).lower()
    if not text or text in {"-", "n/a", "na"}:
        return ""
    if "@" not in text and re.search(r"gmail\.com$", text):
        text = re.sub(r"gmail\.com$", "@gmail.com", text)
    return text[:100]


def parse_students() -> list[dict[str, str]]:
    workbook = openpyxl.load_workbook(WORKBOOK, data_only=True, read_only=True)
    students: list[dict[str, str]] = []
    for sheet_name in workbook.sheetnames:
        class_label = SHEET_TO_CLASS.get(sheet_name.strip().upper())
        if class_label is None:
            continue
        ws = workbook[sheet_name]
        for row in ws.iter_rows(values_only=True):
            cells = [clean(cell) for cell in row]
            number = cells[0] if cells else ""
            name = cells[1] if len(cells) > 1 else ""
            if not number.isdigit() or not name:
                continue
            if name.upper() in {"NAMES", "NAME", "NAME'S", "STUDENT NAME", "STUNDENT'S LIST"}:
                continue
            fname, lname = split_name(name)
            students.append(
                {
                    "class_label": class_label,
                    "sheet": sheet_name.strip(),
                    "full_name": name,
                    "fname": fname,
                    "lname": lname,
                }
            )
    workbook.close()
    return students


def cell_text(cell: ET.Element) -> str:
    return "".join((t.text or "") for t in cell.findall(".//w:t", NS)).strip()


def parse_staff() -> list[dict[str, str]]:
    staff: list[dict[str, str]] = []
    with zipfile.ZipFile(STAFF_DOC) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    for table in root.findall(".//w:tbl", NS):
        started = False
        for row in table.findall("w:tr", NS):
            cells = [clean(cell_text(cell)) for cell in row.findall("w:tc", NS)]
            joined = " ".join(cells).upper()
            if not started:
                if "NAME" in joined and "POSITION" in joined:
                    started = True
                continue
            if not (cells and cells[0].isdigit()):
                continue
            name = cells[1] if len(cells) > 1 else ""
            if not name:
                continue
            fname, lname = split_name(name)
            position = cells[2] if len(cells) > 2 else "TEACHER"
            if position.upper() in {"NOT PROVIDED", ""}:
                position = "TEACHER"
            staff.append(
                {
                    "full_name": name,
                    "fname": fname,
                    "lname": lname,
                    "position": position,
                    "phone": clean_phone(cells[3] if len(cells) > 3 else ""),
                    "email": clean_email(cells[4] if len(cells) > 4 else ""),
                }
            )
    return staff


def parse_photos() -> list[dict[str, str]]:
    photos: list[dict[str, str]] = []
    if not PHOTO_DIR.is_dir():
        return photos
    for path in sorted(PHOTO_DIR.iterdir()):
        if path.suffix.lower() not in {".jpg", ".jpeg", ".png", ".webp"}:
            continue
        photos.append({"file": path.name, "name_hint": path.stem})
    return photos


def main() -> int:
    if not WORKBOOK.is_file():
        raise SystemExit(f"Missing {WORKBOOK}")
    if "BURERA" not in WORKBOOK.name.upper() or "BURERA" not in str(ROOT).upper():
        raise SystemExit("Refusing to parse: workbook is not Wisdom Burera")
    students = parse_students()
    staff = parse_staff()
    photos = parse_photos()
    by_class: dict[str, int] = {}
    for student in students:
        by_class[student["class_label"]] = by_class.get(student["class_label"], 0) + 1
    head = next((row["full_name"] for row in staff if "HEAD" in row["position"].upper()), "")
    payload = {
        "school": {
            "name": "WISDOM SCHOOL BURERA",
            "acronym": "WIS-BUR",
            "slogan": "",
            "academic_year": "2026-2027",
            "term": "Term I",
            "head_teacher": head,
        },
        "classes": students,
        "staff": staff,
        "photos": photos,
        "skipped": [],
    }
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(
        json.dumps(
            {
                "students": len(students),
                "staff": len(staff),
                "photos": len(photos),
                "by_class": by_class,
                "head_teacher": head,
            },
            indent=2,
        )
    )
    print("Staff:", len(staff))
    for row in staff:
        print(" -", row["full_name"], "|", row["position"], "|", row["phone"], "|", row["email"])
    print("Photos:")
    for photo in photos:
        print(" -", photo["file"])
    print("Wrote", OUTPUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
