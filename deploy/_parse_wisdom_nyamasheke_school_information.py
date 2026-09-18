"""Parse Wisdom Nyamasheke students, staff, and named photos.

Scoped to C:\\methode\\15 Wisdoms\\5.Nyamasheke only.
All student sheets are REB: Baby class, Middle Class, Top Class, P1–P6.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(r"C:\methode\15 Wisdoms\5.Nyamasheke")
STUDENTS_XLSX = ROOT / "LIST OF STUDENTS OF WISDOM 2026 - 2027.xlsx"
STAFF_XLSX = ROOT / "teachers identification wisdom school Nyamasheke.xlsx"
OUTPUT = Path(r"C:\xampp7\htdocs\Xander-school\deploy\_wisdom_nyamasheke_school_information.json")

SHEET_TO_CLASS = {
    "N1": "BABY CLASS",
    "N2": "MIDDLE CLASS",
    "N3": "TOP CLASS",
    "P1": "P1",
    "P2": "P2",
    "P3": "P3",
    "P4": "P4",
    "P5": "P5",
    "P6": "P6",
}


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip(" \t:-.")


def split_name(value: str) -> tuple[str, str]:
    parts = clean(value).split(" ")
    if not parts or parts == [""]:
        return "", ""
    if len(parts) == 1:
        return parts[0], ""
    return parts[0], " ".join(parts[1:])


def clean_phone(value: object) -> str:
    digits = re.sub(r"[^0-9]", "", clean(value).lstrip("'"))
    if len(digits) == 9 and digits.startswith("7"):
        digits = "0" + digits
    return digits[:20]


def clean_email(value: object) -> str:
    text = clean(value).lower().replace(" ", "").replace("..", ".")
    if not text or text in {"-", "n/a", "na"}:
        return ""
    text = text.replace("@gnail.com", "@gmail.com")
    if "gmail" in text and "@" not in text:
        text = re.sub(r"gmail\.com$", "@gmail.com", text)
    if "@gmail" in text and not text.endswith("gmail.com"):
        text = re.sub(r"@gmail.*", "@gmail.com", text)
    return text[:100]


def normalize_position(value: str) -> str:
    compact = re.sub(r"[^A-Z]", "", clean(value).upper())
    if "DEPUTY" in compact:
        return "DEPUTY HEAD TEACHER"
    if "HEAD" in compact:
        return "HEAD TEACHER"
    if "DOS" in compact:
        return "DOS"
    if "ACCOUNTANT" in compact or "BURSAR" in compact:
        return "ACCOUNTANT"
    return "TEACHER"


def photo_hint(stem: str) -> str:
    text = re.sub(
        r"\b(P[1-6]|BABY|MIDDLE|TOP|CLSS|CLASS|TEACHER|HEADTEACHER|HEAD|SCHOOL|ACCOUNTANT|DOS)\b",
        " ",
        stem,
        flags=re.I,
    )
    text = re.sub(r"[^A-Za-z ]", " ", text)
    return clean(text)


def parse_students() -> list[dict[str, str]]:
    workbook = openpyxl.load_workbook(STUDENTS_XLSX, data_only=True, read_only=True)
    students: list[dict[str, str]] = []
    seen: set[tuple[str, str]] = set()
    for sheet_name in workbook.sheetnames:
        class_label = SHEET_TO_CLASS.get(sheet_name.strip().upper())
        if class_label is None:
            continue
        ws = workbook[sheet_name]
        for row in ws.iter_rows(values_only=True):
            name = clean(row[0] if row else "")
            if not name or name.upper() in {"NAMES", "NAME", "NO", "N0"}:
                continue
            if not re.search(r"[A-Za-z]", name):
                continue
            key = (class_label, re.sub(r"[^A-Z0-9]", "", name.upper()))
            if key in seen:
                continue
            seen.add(key)
            fname, lname = split_name(name)
            students.append(
                {
                    "class_label": class_label,
                    "sheet": sheet_name.strip(),
                    "full_name": name,
                    "fname": fname,
                    "lname": lname,
                    "sex": "U",
                }
            )
    workbook.close()
    return students


def parse_staff() -> list[dict[str, str]]:
    workbook = openpyxl.load_workbook(STAFF_XLSX, data_only=True, read_only=True)
    ws = workbook[workbook.sheetnames[0]]
    staff: list[dict[str, str]] = []
    started = False
    for row in ws.iter_rows(values_only=True):
        cells = [clean(cell) for cell in (row or ())]
        joined = " ".join(cells).upper()
        if not started:
            if "NAME" in joined and "POST" in joined:
                started = True
            continue
        name = cells[0] if cells else ""
        if not name or not re.search(r"[A-Za-z]", name):
            continue
        if name.upper() in {"NAME", "NAMES"}:
            continue
        fname, lname = split_name(name)
        staff.append(
            {
                "full_name": name,
                "fname": fname,
                "lname": lname,
                "position": normalize_position(cells[1] if len(cells) > 1 else "TEACHER"),
                "phone": clean_phone(cells[2] if len(cells) > 2 else ""),
                "email": clean_email(cells[3] if len(cells) > 3 else ""),
            }
        )
    workbook.close()
    return staff


def parse_photos() -> list[dict[str, str]]:
    photos: list[dict[str, str]] = []
    for path in sorted(ROOT.iterdir()):
        if path.suffix.lower() not in {".jpg", ".jpeg", ".png", ".webp"}:
            continue
        photos.append({"file": path.name, "name_hint": photo_hint(path.stem)})
    return photos


def main() -> int:
    if "NYAMASHEKE" not in str(ROOT).upper():
        raise SystemExit("Refusing to parse: folder is not Wisdom Nyamasheke")
    if not STUDENTS_XLSX.is_file() or not STAFF_XLSX.is_file():
        raise SystemExit("Missing Nyamasheke student or teacher workbook")
    students = parse_students()
    staff = parse_staff()
    photos = parse_photos()
    by_class: dict[str, int] = {}
    for student in students:
        by_class[student["class_label"]] = by_class.get(student["class_label"], 0) + 1
    head = next((row["full_name"] for row in staff if row["position"] == "HEAD TEACHER"), "")
    payload = {
        "school": {
            "name": "WISDOM SCHOOL NYAMASHEKE",
            "acronym": "WIS-NYM",
            "slogan": "",
            "academic_year": "2026-2027",
            "term": "Term I",
            "head_teacher": head,
        },
        "classes": students,
        "staff": staff,
        "photos": photos,
        "skipped": [],
    }
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(
        json.dumps(
            {
                "students": len(students),
                "staff": len(staff),
                "photos": len(photos),
                "by_class": by_class,
                "head_teacher": head,
            },
            indent=2,
        )
    )
    print("Staff:")
    for row in staff:
        print(" -", row["full_name"], "|", row["position"], "|", row["phone"], "|", row["email"])
    print("Photos:")
    for photo in photos:
        print(" -", photo["file"], "=>", photo["name_hint"])
    print("Wrote", OUTPUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
