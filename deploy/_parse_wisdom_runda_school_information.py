"""Parse Wisdom Runda students, staff, and photos.

Scoped to C:\\methode\\15 Wisdoms\\12.Wisdom Runda only.
All student sheets are REB: Baby class, Middle Class, Top Class, P1–P6.
Staff photos are camera files named like TR AIMEEIMG_....jpg.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(r"C:\methode\15 Wisdoms\12.Wisdom Runda")
STUDENTS_XLSX = ROOT / "LIST OF LEARNERS 2026-2027.xlsx"
STAFF_XLSX = ROOT / "TEACHERS' LIST.xlsx"
PHOTO_DIR = ROOT / "IMG_20260918_071518_232"
OUTPUT = Path(r"C:\xampp7\htdocs\Xander-school\deploy\_wisdom_runda_school_information.json")

SHEET_TO_CLASS = {
    "BABY CLASS": "BABY CLASS",
    "BABY": "BABY CLASS",
    "MIDDLE CLASS": "MIDDLE CLASS",
    "MIDDLE": "MIDDLE CLASS",
    "TOP CLASS": "TOP CLASS",
    "TOP": "TOP CLASS",
    "P1": "P1",
    "P2": "P2",
    "P3": "P3",
    "P4": "P4",
    "P5": "P5",
    "P6": "P6",
}


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip(" \t:-.")


def split_name(value: str) -> tuple[str, str]:
    parts = clean(value).split(" ")
    if not parts or parts == [""]:
        return "", ""
    if len(parts) == 1:
        return parts[0], ""
    return parts[0], " ".join(parts[1:])


def clean_phone(value: object) -> str:
    text = clean(value).lstrip("'")
    digits = re.sub(r"[^0-9+]", "", text)
    return digits[:20]


def clean_email(value: object) -> str:
    text = clean(value).lower().replace(" ", "").replace("..", ".")
    if not text or text in {"-", "n/a", "na"}:
        return ""
    text = re.sub(r"@(\d+)gmail\.com$", r"\1@gmail.com", text)
    if "gmail" in text and "@" not in text:
        text = re.sub(r"gmail\.com$", "@gmail.com", text)
    if "@gmail" in text and not text.endswith("gmail.com"):
        text = re.sub(r"@gmail.*", "@gmail.com", text)
    return text[:100]


def clean_sex(value: object) -> str:
    text = clean(value).upper()
    if text in {"M", "MALE", "BOY"}:
        return "M"
    if text in {"F", "FEMALE", "GIRL"}:
        return "F"
    return "U"


def normalize_position(value: str) -> str:
    compact = re.sub(r"[^A-Z]", "", clean(value).upper())
    if "HEAD" in compact:
        return "HEAD TEACHER"
    if "DOS" in compact:
        return "DOS"
    if "ACCOUNTANT" in compact:
        return "ACCOUNTANT"
    return "TEACHER"


def photo_hint(stem: str) -> str:
    text = re.sub(r"^TR\s*", "", stem, flags=re.I)
    text = re.split(r"IMG[_-]?\d", text, maxsplit=1, flags=re.I)[0]
    text = re.sub(r"[^A-Za-z ]", " ", text)
    return clean(text)


def parse_students() -> list[dict[str, str]]:
    workbook = openpyxl.load_workbook(STUDENTS_XLSX, data_only=True, read_only=True)
    students: list[dict[str, str]] = []
    for sheet_name in workbook.sheetnames:
        class_label = SHEET_TO_CLASS.get(sheet_name.strip().upper())
        if class_label is None:
            continue
        ws = workbook[sheet_name]
        for row in ws.iter_rows(values_only=True):
            cells = [clean(cell) for cell in row]
            number = cells[0] if cells else ""
            name = cells[1] if len(cells) > 1 else ""
            sex = clean_sex(cells[2] if len(cells) > 2 else "")
            if not number.isdigit() or not name:
                continue
            if name.upper() in {"STUDENT NAME", "NAMES", "NAME"}:
                continue
            fname, lname = split_name(name)
            students.append(
                {
                    "class_label": class_label,
                    "sheet": sheet_name.strip(),
                    "full_name": name,
                    "fname": fname,
                    "lname": lname,
                    "sex": sex,
                }
            )
    workbook.close()
    return students


def parse_staff() -> list[dict[str, str]]:
    workbook = openpyxl.load_workbook(STAFF_XLSX, data_only=True, read_only=True)
    ws = workbook[workbook.sheetnames[0]]
    staff: list[dict[str, str]] = []
    started = False
    for row in ws.iter_rows(values_only=True):
        cells = [clean(cell) for cell in row]
        joined = " ".join(cells).upper()
        if not started:
            if "NAMES" in joined and "POSITION" in joined:
                started = True
            continue
        number = ""
        name = ""
        position = "TEACHER"
        phone = ""
        email = ""
        if cells and cells[0].isdigit():
            number, name = cells[0], cells[1] if len(cells) > 1 else ""
            position = cells[2] if len(cells) > 2 else "TEACHER"
            phone = cells[3] if len(cells) > 3 else ""
            email = cells[4] if len(cells) > 4 else ""
        elif len(cells) > 1 and cells[1].isdigit():
            number, name = cells[1], cells[2] if len(cells) > 2 else ""
            position = cells[3] if len(cells) > 3 else "TEACHER"
            phone = cells[4] if len(cells) > 4 else ""
            email = cells[5] if len(cells) > 5 else ""
        if not number or not name:
            continue
        fname, lname = split_name(name)
        staff.append(
            {
                "full_name": name,
                "fname": fname,
                "lname": lname,
                "position": normalize_position(position),
                "phone": clean_phone(phone),
                "email": clean_email(email),
            }
        )
    workbook.close()
    return staff


def parse_photos() -> list[dict[str, str]]:
    photos: list[dict[str, str]] = []
    if not PHOTO_DIR.is_dir():
        return photos
    for path in sorted(PHOTO_DIR.iterdir()):
        if path.suffix.lower() not in {".jpg", ".jpeg", ".png", ".webp"}:
            continue
        photos.append({"file": path.name, "name_hint": photo_hint(path.stem)})
    return photos


def main() -> int:
    if not STUDENTS_XLSX.is_file() or not STAFF_XLSX.is_file():
        raise SystemExit("Missing Runda student or teacher workbook")
    if "RUNDA" not in str(ROOT).upper():
        raise SystemExit("Refusing to parse: folder is not Wisdom Runda")
    students = parse_students()
    staff = parse_staff()
    photos = parse_photos()
    by_class: dict[str, int] = {}
    for student in students:
        by_class[student["class_label"]] = by_class.get(student["class_label"], 0) + 1
    head = next((row["full_name"] for row in staff if "HEAD" in row["position"].upper()), "")
    payload = {
        "school": {
            "name": "WISDOM SCHOOL RUNDA",
            "acronym": "WIS-RUN",
            "slogan": "",
            "academic_year": "2026-2027",
            "term": "Term I",
            "head_teacher": head,
        },
        "classes": students,
        "staff": staff,
        "photos": photos,
        "skipped": [],
    }
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(
        json.dumps(
            {
                "students": len(students),
                "staff": len(staff),
                "photos": len(photos),
                "by_class": by_class,
                "head_teacher": head,
            },
            indent=2,
        )
    )
    print("Staff:")
    for row in staff:
        print(" -", row["full_name"], "|", row["position"], "|", row["phone"], "|", row["email"])
    print("Photos:")
    for photo in photos:
        print(" -", photo["file"], "=>", photo["name_hint"])
    print("Wrote", OUTPUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
