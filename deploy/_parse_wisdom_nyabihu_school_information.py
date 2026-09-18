"""Parse Wisdom Nyabihu students, staff, and photos.

Scoped to C:\\methode\\15 Wisdoms\\10.Wisdom NYABIHU only.
All student sheets are REB: Baby class, Middle Class, Top Class, P1–P6.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(r"C:\methode\15 Wisdoms\10.Wisdom NYABIHU")
WORKBOOK = ROOT / "REQUESTED REPORT NYABIHU" / "requested report WS NYABIHU.xlsx"
PHOTO_DIR = ROOT / "REQUESTED REPORT NYABIHU"
OUTPUT = Path(r"C:\xampp7\htdocs\Xander-school\deploy\_wisdom_nyabihu_school_information.json")

SHEET_TO_CLASS = {
    "BABY": "BABY CLASS",
    "MIDDLE": "MIDDLE CLASS",
    "TOP": "TOP CLASS",
    "P1": "P1",
    "P2": "P2",
    "P3": "P3",
    "P4": "P4",
    "P5": "P5",
    "P6": "P6",
}


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def split_name(value: str) -> tuple[str, str]:
    parts = clean(value).split(" ")
    if not parts or parts == [""]:
        return "", ""
    if len(parts) == 1:
        return parts[0], ""
    return parts[0], " ".join(parts[1:])


def compact(value: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", value.upper())


def clean_phone(value: object) -> str:
    text = clean(value).lstrip("'")
    digits = re.sub(r"[^0-9+]", "", text)
    return digits[:20]


def parse_students() -> list[dict[str, str]]:
    workbook = openpyxl.load_workbook(WORKBOOK, data_only=True, read_only=True)
    students: list[dict[str, str]] = []
    for sheet_name in workbook.sheetnames:
        if sheet_name.upper() == "TEACHERS":
            continue
        class_label = SHEET_TO_CLASS.get(sheet_name.upper())
        if class_label is None:
            continue
        ws = workbook[sheet_name]
        rows = [list(row) for row in ws.iter_rows(values_only=True)]
        header_blob = " ".join(clean(cell) for row in rows[:6] for cell in row[:3]).upper()
        if "NYABIHU" not in header_blob:
            continue
        for row in rows:
            cells = [clean(cell) for cell in row]
            number = cells[0] if cells else ""
            name = cells[1] if len(cells) > 1 else ""
            if not number.isdigit() or not name:
                continue
            if name.upper() in {"NAMES", "NAME", "STUDENT NAME", "STUNDENT'S LIST", "STUNDENT'SLIST"}:
                continue
            fname, lname = split_name(name)
            students.append(
                {
                    "class_label": class_label,
                    "sheet": sheet_name,
                    "full_name": name,
                    "fname": fname,
                    "lname": lname,
                }
            )
    workbook.close()
    return students


def parse_staff() -> list[dict[str, str]]:
    workbook = openpyxl.load_workbook(WORKBOOK, data_only=True, read_only=True)
    ws = workbook["TEACHERS"]
    staff: list[dict[str, str]] = []
    started = False
    for row in ws.iter_rows(values_only=True):
        cells = [clean(cell) for cell in row]
        joined = " ".join(cells).upper()
        if not started:
            if "NAMES" in joined and "POSITION" in joined:
                started = True
            continue
        if not (cells and cells[0].isdigit()):
            continue
        name = cells[1] if len(cells) > 1 else ""
        if not name:
            continue
        fname, lname = split_name(name)
        staff.append(
            {
                "full_name": name,
                "fname": fname,
                "lname": lname,
                "position": cells[2] if len(cells) > 2 else "TEACHER",
                "phone": clean_phone(cells[3] if len(cells) > 3 else ""),
                "email": clean(cells[4] if len(cells) > 4 else "").lower(),
            }
        )
    workbook.close()
    return staff


def parse_photos() -> list[dict[str, str]]:
    photos: list[dict[str, str]] = []
    if not PHOTO_DIR.is_dir():
        return photos
    for path in sorted(PHOTO_DIR.iterdir()):
        if path.suffix.lower() not in {".jpg", ".jpeg", ".png", ".webp"}:
            continue
        photos.append({"file": path.name, "name_hint": path.stem})
    return photos


def main() -> int:
    if not WORKBOOK.is_file():
        raise SystemExit(f"Missing {WORKBOOK}")
    students = parse_students()
    staff = parse_staff()
    photos = parse_photos()
    by_class: dict[str, int] = {}
    for student in students:
        by_class[student["class_label"]] = by_class.get(student["class_label"], 0) + 1
    head = next((row["full_name"] for row in staff if "HEAD" in row["position"].upper()), "")
    payload = {
        "school": {
            "name": "WISDOM SCHOOL NYABIHU",
            "acronym": "WIS-NYB",
            "slogan": "",
            "academic_year": "2026-2027",
            "term": "Term I",
            "head_teacher": head,
        },
        "classes": students,
        "staff": staff,
        "photos": photos,
        "skipped": [],
    }
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(
        json.dumps(
            {
                "students": len(students),
                "staff": len(staff),
                "photos": len(photos),
                "by_class": by_class,
            },
            indent=2,
        )
    )
    print("Staff:", len(staff))
    for row in staff:
        print(" -", row["full_name"], "|", row["position"], "|", row["phone"])
    print("Wrote", OUTPUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
