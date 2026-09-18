"""Parse Wisdom Ngororero students, staff, and photos.

Scoped to C:\\methode\\15 Wisdoms\\1.Wisdom Ngororero only.
All student sheets are REB: Baby class, Middle Class, Top Class, P1–P6.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(r"C:\methode\15 Wisdoms\1.Wisdom Ngororero")
WORKBOOK = ROOT / "Students and Staffs.xlsx"
PHOTO_DIR = ROOT / "Staff photo"
OUTPUT = Path(r"C:\xampp7\htdocs\Xander-school\deploy\_wisdom_ngororero_school_information.json")

SHEET_TO_CLASS = {
    "BABY": "BABY CLASS",
    "BABY CLASS": "BABY CLASS",
    "MIDDLE": "MIDDLE CLASS",
    "MIDDLE CLASS": "MIDDLE CLASS",
    "TOP": "TOP CLASS",
    "TOP CLASS": "TOP CLASS",
    "P1": "P1",
    "P2": "P2",
    "P3": "P3",
    "P4": "P4",
    "P5": "P5",
    "P6": "P6",
}
NAME_HEADERS = {
    "NAMES",
    "NAME",
    "NAME'S",
    "STUDENT NAME",
    "STUNDENT'S LIST",
}


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip(" \t:-")


def split_name(value: str) -> tuple[str, str]:
    parts = clean(value).split(" ")
    if not parts or parts == [""]:
        return "", ""
    if len(parts) == 1:
        return parts[0], ""
    return parts[0], " ".join(parts[1:])


def clean_phone(value: object) -> str:
    text = clean(value).lstrip("'")
    digits = re.sub(r"[^0-9+]", "", text)
    return digits[:20]


def clean_email(value: object) -> str:
    text = clean(value).lower().replace(" ", "")
    if not text or text in {"-", "n/a", "na"}:
        return ""
    if "@" not in text and re.search(r"gmail\.com$", text):
        text = re.sub(r"gmail\.com$", "@gmail.com", text)
    return text[:100]


def student_name(cells: list[str]) -> str:
    parts: list[str] = []
    for cell in cells[1:4]:
        text = clean(cell)
        if not text or text.upper() in NAME_HEADERS:
            continue
        if re.fullmatch(r"[0-9.]+", text):
            continue
        parts.append(text)
    return clean(" ".join(parts))


def normalize_position(value: str) -> str:
    text = clean(value)
    compact = re.sub(r"[^A-Z]", "", text.upper())
    if "HEAD" in compact:
        return "HEAD TEACHER"
    if "DOS" in compact:
        return "DOS"
    if "ACCOUNTANT" in compact:
        return "ACCOUNTANT"
    if not text or text.upper() in {"NOT PROVIDED", "TEACHER.", "TEACHER", "TEEEACHER"}:
        return "TEACHER"
    return text


def parse_students() -> list[dict[str, str]]:
    workbook = openpyxl.load_workbook(WORKBOOK, data_only=True, read_only=True)
    students: list[dict[str, str]] = []
    for sheet_name in workbook.sheetnames:
        class_label = SHEET_TO_CLASS.get(sheet_name.strip().upper())
        if class_label is None:
            continue
        ws = workbook[sheet_name]
        for row in ws.iter_rows(values_only=True):
            cells = [clean(cell) for cell in row]
            number = cells[0] if cells else ""
            if not number.isdigit():
                continue
            name = student_name(cells)
            if not name or name.upper() in NAME_HEADERS:
                continue
            fname, lname = split_name(name)
            students.append(
                {
                    "class_label": class_label,
                    "sheet": sheet_name.strip(),
                    "full_name": name,
                    "fname": fname,
                    "lname": lname,
                }
            )
    workbook.close()
    return students


def parse_staff() -> list[dict[str, str]]:
    workbook = openpyxl.load_workbook(WORKBOOK, data_only=True, read_only=True)
    staff: list[dict[str, str]] = []
    for sheet_name in workbook.sheetnames:
        if "STAFF" not in sheet_name.upper():
            continue
        ws = workbook[sheet_name]
        started = False
        for row in ws.iter_rows(values_only=True):
            cells = [clean(cell) for cell in row]
            joined = " ".join(cells).upper()
            if not started:
                if "NAME" in joined and "POSITION" in joined:
                    started = True
                continue
            name = cells[0] if cells else ""
            if not name or name.upper() in NAME_HEADERS:
                continue
            fname, lname = split_name(name)
            staff.append(
                {
                    "full_name": name,
                    "fname": fname,
                    "lname": lname,
                    "position": normalize_position(cells[1] if len(cells) > 1 else "TEACHER"),
                    "phone": clean_phone(cells[2] if len(cells) > 2 else ""),
                    "email": clean_email(cells[3] if len(cells) > 3 else ""),
                }
            )
    workbook.close()
    return staff


def parse_photos() -> list[dict[str, str]]:
    photos: list[dict[str, str]] = []
    if not PHOTO_DIR.is_dir():
        return photos
    for path in sorted(PHOTO_DIR.iterdir()):
        if path.suffix.lower() not in {".jpg", ".jpeg", ".png", ".webp"}:
            continue
        photos.append({"file": path.name, "name_hint": path.stem})
    return photos


def main() -> int:
    if not WORKBOOK.is_file():
        raise SystemExit(f"Missing {WORKBOOK}")
    if "NGORORERO" not in str(ROOT).upper():
        raise SystemExit("Refusing to parse: folder is not Wisdom Ngororero")
    students = parse_students()
    staff = parse_staff()
    photos = parse_photos()
    by_class: dict[str, int] = {}
    for student in students:
        by_class[student["class_label"]] = by_class.get(student["class_label"], 0) + 1
    head = next((row["full_name"] for row in staff if "HEAD" in row["position"].upper()), "")
    payload = {
        "school": {
            "name": "WISDOM SCHOOL NGORORERO",
            "acronym": "WIS-NGO",
            "slogan": "",
            "academic_year": "2026-2027",
            "term": "Term I",
            "head_teacher": head,
        },
        "classes": students,
        "staff": staff,
        "photos": photos,
        "skipped": [],
    }
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(
        json.dumps(
            {
                "students": len(students),
                "staff": len(staff),
                "photos": len(photos),
                "by_class": by_class,
                "head_teacher": head,
            },
            indent=2,
        )
    )
    print("Staff:", len(staff))
    for row in staff:
        print(" -", row["full_name"], "|", row["position"], "|", row["phone"], "|", row["email"])
    print("Photos:")
    for photo in photos:
        print(" -", photo["file"])
    print("Wrote", OUTPUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
